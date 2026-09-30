import csv
from collections import Counter
from pathlib import Path

import openpyxl


IDENT_COLUMNS = ["no", "layer", "name", "domain", "type"]
MASTER_SHEET = "마스터_표본프레임"
SUCCESS = {"ok", "not_found"}


def _read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        fields = reader.fieldnames or []
        rows = list(reader)
    if len(fields) != len(set(fields)):
        raise ValueError("Duplicate CSV columns: " + str(path))
    required = IDENT_COLUMNS + ["fetch_status"]
    if any(column not in fields for column in required):
        raise ValueError("Missing required CSV columns: " + str(path))
    if any(None in row or any(value is None for value in row.values()) for row in rows):
        raise ValueError("Malformed CSV record: " + str(path))
    return fields, rows


def _index(rows, label):
    result = {}
    for row in rows:
        domain = row["domain"].strip().lower()
        if not domain or domain in result:
            raise ValueError("Empty or duplicate domain in " + label + ": " + domain)
        result[domain] = row
    return result


def reconstruct(master_path: Path, initial_path: Path, retry_path: Path, out_dir: Path) -> dict:
    master_path, initial_path, retry_path, out_dir = map(
        Path, (master_path, initial_path, retry_path, out_dir)
    )
    workbook = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    master_rows = []
    source_rows = excluded_n = missing_domain_rows = 0
    try:
        if MASTER_SHEET not in workbook.sheetnames:
            raise ValueError("Missing master sheet: " + MASTER_SHEET)
        sheet = workbook[MASTER_SHEET]
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        expected = {0: "No", 1: "층", 2: "명칭", 3: "도메인", 4: "유형", 8: "수집대상"}
        if any(len(header) <= index or header[index] != name for index, name in expected.items()):
            raise ValueError("Master columns differ from the archived collection schema")
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if values[0] is None:
                continue
            source_rows += 1
            target = str(values[8]).strip().upper() if values[8] is not None else ""
            if target == "N":
                excluded_n += 1
                continue
            domain = str(values[3]).strip().lower() if values[3] else None
            if not domain:
                missing_domain_rows += 1
                continue
            master_rows.append(
                {
                    "no": str(int(values[0])),
                    "layer": str(values[1]).strip() if values[1] else "",
                    "name": str(values[2]).strip() if values[2] else "",
                    "domain": domain,
                    "type": str(values[4]).strip() if values[4] else "",
                }
            )
    finally:
        workbook.close()
    master_unique = {}
    for row in master_rows:
        master_unique.setdefault(row["domain"], row)
    fields, initial_rows = _read_csv(initial_path)
    retry_fields, retry_rows = _read_csv(retry_path)
    if retry_fields != fields:
        raise ValueError("Initial and retry CSV schemas differ")
    initial = _index(initial_rows, "initial snapshot")
    retries = _index(retry_rows, "retry snapshot")
    missing_initial = sorted(set(master_unique) - set(initial))
    unused_initial = sorted(set(initial) - set(master_unique))
    if missing_initial or unused_initial:
        raise ValueError(
            "Master/initial domain mismatch: missing="
            + repr(missing_initial)
            + "; unexpected="
            + repr(unused_initial)
        )
    unknown_retries = sorted(set(retries) - set(initial))
    if unknown_retries:
        raise ValueError(
            "Retry domains are missing from the initial snapshot: " + repr(unknown_retries)
        )
    merged = dict(initial)
    recovered = []
    for domain, row in retries.items():
        if row["fetch_status"] in SUCCESS and merged[domain]["fetch_status"] not in SUCCESS:
            merged[domain] = row
            recovered.append(domain)
    rebuilt = [
        {
            column: row[column] if column in IDENT_COLUMNS else merged[row["domain"]][column]
            for column in fields
        }
        for row in master_rows
    ]
    destination = out_dir / "snapshot_rebuilt.csv"
    if destination.resolve() in {
        path.resolve() for path in (master_path, initial_path, retry_path)
    }:
        raise ValueError("Output must not overwrite an input")
    out_dir.mkdir(parents=True, exist_ok=True)
    with destination.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rebuilt)
    return {
        "master_source_rows": source_rows,
        "excluded_target_n_rows": excluded_n,
        "excluded_missing_domain_rows": missing_domain_rows,
        "master_selected_rows": len(master_rows),
        "master_unique_domains": len(master_unique),
        "duplicate_rows_removed_for_collection": len(master_rows) - len(master_unique),
        "initial_domains": len(initial),
        "initial_status_counts": dict(Counter(row["fetch_status"] for row in initial_rows)),
        "retry_domains": len(retries),
        "retry_status_counts": dict(Counter(row["fetch_status"] for row in retry_rows)),
        "recovered_domains": recovered,
        "recovered_domain_count": len(recovered),
        "final_unique_status_counts": dict(Counter(row["fetch_status"] for row in merged.values())),
        "rebuilt_rows": len(rebuilt),
        "rebuilt_status_counts": dict(Counter(row["fetch_status"] for row in rebuilt)),
        "rebuilt_columns": fields,
    }
