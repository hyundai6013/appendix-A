import argparse
import csv
import os
import random
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

INPUTS = Path(__file__).resolve().parents[1] / "data" / "inputs"


try:
    import requests
except ImportError:
    sys.exit("[의존성 오류] 'requests' 미설치.  ->  pip install requests protego openpyxl")

try:
    from protego import Protego
except ImportError:
    sys.exit(
        "[의존성 오류] 'protego' 미설치.  ->  pip install protego\n"
        "  (RFC 9309 준수 파서. urllib.robotparser는 요구사항상 지양)"
    )

try:
    import openpyxl
except ImportError:
    sys.exit("[의존성 오류] 'openpyxl' 미설치.  ->  pip install openpyxl")


BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/126.0.0.0 Safari/537.36"
)


NEUTRAL_UA = "Mozilla/5.0 (compatible; UnlistedGenericBot/1.0)"


PROBE_PATH = "/"


def _header_map(ws):

    headers = {}
    try:
        first_row = next(ws.iter_rows(min_row=1, max_row=1))
    except StopIteration:
        return headers
    for j, cell in enumerate(first_row):
        if cell.value is not None:
            headers[str(cell.value).strip()] = j
    return headers


def load_crawlers(ws):

    hmap = _header_map(ws)
    ua_col = hmap.get("크롤러 User-Agent", 2)
    no_col = hmap.get("No", 0)
    crawlers = []
    for row in ws.iter_rows(min_row=2):
        no_val = row[no_col].value if no_col < len(row) else None
        if not isinstance(no_val, (int, float)):
            continue
        ua = row[ua_col].value if ua_col < len(row) else None
        if ua and str(ua).strip():
            crawlers.append(str(ua).strip())
    return crawlers


def load_sites(ws, layer_hint=""):

    hmap = _header_map(ws)
    no_col = hmap.get("No", 0)
    layer_col = hmap.get("층", 1)
    domain_col = hmap.get("도메인", 3)

    name_col = next((hmap[k] for k in ("매체명", "기관명", "서비스명", "명칭") if k in hmap), 2)

    type_col = next((v for k, v in hmap.items() if k.endswith("유형")), 4)

    target_col = hmap.get("수집대상", None)

    sites = []
    for row in ws.iter_rows(min_row=2):
        no_val = row[no_col].value if no_col < len(row) else None
        if not isinstance(no_val, (int, float)):
            continue
        domain = row[domain_col].value if domain_col < len(row) else None
        if not domain or not str(domain).strip():
            continue
        if target_col is not None and target_col < len(row) and row[target_col].value is not None:
            collect_target = str(row[target_col].value).strip()
        else:
            collect_target = ""
        sites.append(
            {
                "no": int(no_val),
                "layer": str(row[layer_col].value).strip()
                if layer_col < len(row) and row[layer_col].value
                else layer_hint,
                "name": str(row[name_col].value).strip()
                if name_col < len(row) and row[name_col].value
                else "",
                "domain": str(domain).strip().lower(),
                "type": str(row[type_col].value).strip()
                if type_col < len(row) and row[type_col].value
                else "",
                "collect_target": collect_target,
            }
        )
    return sites


def load_inputs(xlsx_path, sheet_filter=None, crawlers_xlsx=None):

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    crawlers = []
    sites = []
    for ws in wb.worksheets:
        title = ws.title
        if "크롤러" in title:
            crawlers = load_crawlers(ws)
        elif "표본" in title:
            if sheet_filter and sheet_filter not in title:
                continue

            hint = title.split("_")[-1] if "_" in title else title
            sites.extend(load_sites(ws, layer_hint=hint))
    wb.close()

    if not crawlers and crawlers_xlsx and os.path.exists(crawlers_xlsx):
        cwb = openpyxl.load_workbook(crawlers_xlsx, read_only=True, data_only=True)
        for ws in cwb.worksheets:
            if "크롤러" in ws.title:
                crawlers = load_crawlers(ws)
                break
        cwb.close()
    return crawlers, sites


def build_candidate_hosts(domain, www_fallback):

    hosts = [domain]
    if not www_fallback or domain.startswith("www."):
        return hosts
    labels = domain.split(".")
    kr_second = {"co", "or", "go", "ac", "re", "ne", "kr", "pe"}

    is_bare = (len(labels) == 2) or (len(labels) == 3 and labels[1] in kr_second)
    if is_bare:
        hosts.append("www." + domain)
    return hosts


def fetch_robots(session, domain, timeout, www_fallback):

    result = {
        "fetch_status": "other_error",
        "http_status": None,
        "final_url": None,
        "content": None,
        "error": "",
        "elapsed_sec": 0.0,
    }
    hosts = build_candidate_hosts(domain, www_fallback)
    t0 = time.time()
    last_err = ""
    for host in hosts:
        url = "https://{}/robots.txt".format(host)
        try:
            resp = session.get(url, timeout=timeout, allow_redirects=True)
            result["http_status"] = resp.status_code
            result["final_url"] = resp.url
            if resp.status_code == 200:
                result["content"] = resp.content.decode("utf-8", errors="replace")
                result["fetch_status"] = "ok"
            elif resp.status_code in (404, 410):
                result["fetch_status"] = "not_found"
                result["content"] = ""
            else:
                result["fetch_status"] = "http_error"
                result["error"] = "HTTP {}".format(resp.status_code)
            break
        except requests.exceptions.Timeout as e:
            result["fetch_status"] = "timeout"
            last_err = "timeout: " + str(e)
        except requests.exceptions.SSLError as e:
            result["fetch_status"] = "ssl_error"
            last_err = "ssl: " + str(e)
        except requests.exceptions.TooManyRedirects as e:
            result["fetch_status"] = "too_many_redirects"
            last_err = "redirects: " + str(e)
        except requests.exceptions.ConnectionError as e:
            result["fetch_status"] = "connection_error"
            last_err = "conn: " + str(e)
        except requests.exceptions.RequestException as e:
            result["fetch_status"] = "other_error"
            last_err = "req: " + str(e)

    if result["fetch_status"] not in ("ok", "not_found", "http_error"):
        result["error"] = last_err
    result["elapsed_sec"] = round(time.time() - t0, 3)
    return result


def parse_useragent_tokens(content):

    tokens = set()
    for line in content.splitlines():
        s = line.split("#", 1)[0].strip()
        if not s or s.startswith("#"):
            continue
        low = s.lower()
        if low.startswith("user-agent:"):
            val = s.split(":", 1)[1].strip()
            if val:
                tokens.add(val.lower())
    return tokens


def analyze(content, crawlers, domain):

    rp = Protego.parse(content or "")
    base = "https://{}{}".format(domain, PROBE_PATH)

    decisions = {}
    explicit = {}
    named = parse_useragent_tokens(content or "")
    for c in crawlers:
        allowed = rp.can_fetch(base, c)
        decisions[c] = "ALLOW" if allowed else "BLOCK"
        explicit[c] = c.lower() in named

    global_block_all = not rp.can_fetch(base, NEUTRAL_UA)

    try:
        cd = rp.crawl_delay(NEUTRAL_UA)
    except Exception:
        cd = None

    try:
        sitemap_count = sum(1 for _ in rp.sitemaps)
    except Exception:
        sitemap_count = 0

    return {
        "decisions": decisions,
        "explicit": explicit,
        "global_block_all": global_block_all,
        "crawl_delay_wildcard": cd,
        "sitemap_count": sitemap_count,
    }


def save_raw(raw_dir, domain, content):

    os.makedirs(raw_dir, exist_ok=True)
    safe = domain.replace("/", "_").replace(":", "_")
    path = os.path.join(raw_dir, safe + ".txt")
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write(content if content is not None else "")


def dedup_sites(sites):

    seen = set()
    uniq = []
    for s in sites:
        d = s["domain"]
        if d in seen:
            continue
        seen.add(d)
        uniq.append(s)
    return uniq, len(sites) - len(uniq)


def filter_target(sites, include_nontarget):

    if include_nontarget:
        return sites, 0
    kept = [s for s in sites if str(s.get("collect_target", "")).strip().upper() != "N"]
    return kept, len(sites) - len(kept)


def cap_per_layer(sites, n):

    cap = Counter()
    kept = []
    for s in sites:
        L = s["layer"]
        if cap[L] < n:
            cap[L] += 1
            kept.append(s)
    return kept


def write_summary(wide_path, long_path, raw_dir, summary_path, stamp, title="수집"):

    rows = []
    with open(wide_path, "r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    n = len(rows)
    status_counter = Counter(r.get("fetch_status", "") for r in rows)
    ok = status_counter.get("ok", 0)
    nf = status_counter.get("not_found", 0)
    success = ok + nf

    layer_total = Counter()
    layer_success = Counter()
    layer_block = Counter()
    for r in rows:
        L = r.get("layer", "") or "(미지정)"
        layer_total[L] += 1
        if r.get("fetch_status") in ("ok", "not_found"):
            layer_success[L] += 1
        if str(r.get("global_block_all", "")).strip() in ("True", "TRUE", "1"):
            layer_block[L] += 1

    lines = []
    lines.append("=== {} 요약 ({}) ===".format(title, stamp))
    lines.append("대상 도메인(고유): {}".format(n))
    lines.append(
        "응답 수신(성공): {}/{}  = {:.1f}%".format(success, n, 100.0 * success / n if n else 0)
    )
    lines.append("  - robots.txt 존재(200): {}".format(ok))
    lines.append("  - robots.txt 부재(404/410): {}".format(nf))
    lines.append("")
    lines.append("상태별 분포:")
    for k in sorted(status_counter):
        lines.append("    {:<20} {}".format(k, status_counter[k]))
    lines.append("")
    lines.append("층별 건수 / 응답성공 / 전체차단(global_block_all):")
    for L in sorted(layer_total):
        lines.append(
            "    {:<16} 건수 {:>4}  성공 {:>4} ({:.0f}%)  전체차단 {:>3}".format(
                L,
                layer_total[L],
                layer_success[L],
                100.0 * layer_success[L] / layer_total[L] if layer_total[L] else 0,
                layer_block[L],
            )
        )
    lines.append("")
    lines.append("출력:")
    lines.append("  매트릭스 CSV : {}".format(wide_path))
    lines.append("  tidy long CSV: {}".format(long_path))
    lines.append("  원본 robots  : {}/".format(raw_dir))
    summary = "\n".join(lines)
    with open(summary_path, "w", encoding="utf-8", newline="") as f:
        f.write(summary + "\n")
    return summary


def main():

    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="robots.txt 수집기 (RFC 9309 / protego 기반)")
    ap.add_argument(
        "--xlsx", default=str(INPUTS / "master_frame.xlsx"), help="입력 엑셀 경로(도메인 시트)"
    )
    ap.add_argument(
        "--crawlers-xlsx",
        default=str(INPUTS / "crawler_frame.xlsx"),
        help="크롤러 시트가 있는 엑셀(본 파일에 크롤러 시트가 없을 때 승계)",
    )
    ap.add_argument(
        "--pilot",
        action="store_true",
        help="파일럿 모드: 앞에서 100개 도메인만 (--limit 미지정 시)",
    )
    ap.add_argument("--limit", type=int, default=None, help="처리할 도메인 수 상한 (앞에서 N개)")
    ap.add_argument(
        "--per-layer", type=int, default=None, help="각 층마다 앞에서 N개씩만 (소규모 검증용)"
    )
    ap.add_argument("--start", type=int, default=0, help="건너뛸 도메인 수(재개용 오프셋)")
    ap.add_argument(
        "--sheets",
        default=None,
        help="특정 층만: 시트명 부분문자열 (예: 언론 / 공공교육 / 콘텐츠플랫폼)",
    )
    ap.add_argument("--timeout", type=float, default=10.0, help="요청 타임아웃(초)")
    ap.add_argument("--delay-min", type=float, default=1.0, help="요청 간 최소 대기(초)")
    ap.add_argument("--delay-max", type=float, default=2.0, help="요청 간 최대 대기(초)")
    ap.add_argument(
        "--www-fallback", action="store_true", help="연결 실패 시 www. 접두 호스트로 1회 재시도"
    )
    ap.add_argument(
        "--no-dedup", action="store_true", help="도메인 중복 제거 비활성화(기본은 중복 제거)"
    )
    ap.add_argument(
        "--include-nontarget", action="store_true", help="수집대상=N 행도 포함(기본은 제외)"
    )
    ap.add_argument("--outdir", required=True, help="새 실행의 출력 폴더(기존 폴더 사용 불가)")
    args = ap.parse_args()
    if args.timeout <= 0 or args.delay_min < 0 or args.delay_max < args.delay_min:
        ap.error("timeout > 0 및 0 <= delay-min <= delay-max 조건이 필요합니다.")
    if any(v is not None and v < 0 for v in (args.start, args.limit, args.per_layer)):
        ap.error("start, limit, per-layer는 음수일 수 없습니다.")
    if Path(args.outdir).exists():
        ap.error("출력 경로가 이미 존재합니다. 새로운 실행 폴더를 지정하세요.")

    if not os.path.exists(args.xlsx):
        sys.exit("[오류] 엑셀 파일을 찾을 수 없음: {}".format(args.xlsx))

    crawlers, sites = load_inputs(
        args.xlsx, sheet_filter=args.sheets, crawlers_xlsx=args.crawlers_xlsx
    )
    if not crawlers:
        sys.exit("[오류] 크롤러 목록을 읽지 못함(--crawlers-xlsx 의 시트명에 '크롤러' 포함 확인).")
    if not sites:
        sys.exit("[오류] 도메인 목록을 읽지 못함(시트명에 '표본' 포함 확인).")

    n_raw = len(sites)

    sites, excluded_n = filter_target(sites, args.include_nontarget)

    if args.no_dedup:
        dedup_removed = 0
    else:
        sites, dedup_removed = dedup_sites(sites)

    if args.per_layer is not None:
        sites = cap_per_layer(sites, args.per_layer)

    limit = args.limit
    if limit is None and args.pilot:
        limit = 100
    sel = sites[args.start :]
    if limit is not None:
        sel = sel[:limit]

    print(
        "입력 {}행 -> 수집대상=N 제외 {} / 중복제거 {} -> 고유 {}개".format(
            n_raw, excluded_n, dedup_removed, len(sites)
        )
    )
    print("크롤러 {}종, 이번 실행 대상 {}개".format(len(crawlers), len(sel)))
    print("크롤러:", ", ".join(crawlers))

    os.makedirs(args.outdir, exist_ok=False)
    raw_dir = os.path.join(args.outdir, "robots_raw")

    base_cols = [
        "no",
        "layer",
        "name",
        "domain",
        "type",
        "fetched_at",
        "fetch_status",
        "http_status",
        "final_url",
        "robots_exists",
        "robots_bytes",
        "global_block_all",
        "crawl_delay_wildcard",
        "sitemap_count",
        "elapsed_sec",
        "error_message",
    ]
    block_cols = ["block__" + c for c in crawlers]
    expl_cols = ["explicit__" + c for c in crawlers]
    wide_cols = base_cols + block_cols + expl_cols

    completed = set()
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    wide_path = os.path.join(args.outdir, "robots_results_{}.csv".format(stamp))
    long_path = os.path.join(args.outdir, "robots_matrix_long_{}.csv".format(stamp))
    file_mode = "w"
    summary_path = os.path.join(args.outdir, "run_summary_{}.txt".format(stamp))

    session = requests.Session()
    session.headers.update({"User-Agent": BROWSER_UA, "Accept": "text/plain,*/*"})

    todo = [s for s in sel if s["domain"] not in completed]
    n = len(todo)
    if n == 0:
        print("이번 실행에서 새로 수집할 도메인이 없습니다(모두 완료됨).")

    with (
        open(wide_path, file_mode, encoding="utf-8-sig", newline="") as wf,
        open(long_path, file_mode, encoding="utf-8-sig", newline="") as lf,
    ):
        ww = csv.DictWriter(wf, fieldnames=wide_cols)
        lw = csv.writer(lf)
        if file_mode == "w":
            ww.writeheader()
            lw.writerow(
                [
                    "domain",
                    "layer",
                    "name",
                    "type",
                    "crawler",
                    "decision",
                    "blocked",
                    "explicitly_named",
                    "robots_exists",
                    "fetch_status",
                ]
            )

        for i, site in enumerate(todo, 1):
            domain = site["domain"]
            fetched_at = datetime.now().isoformat(timespec="seconds")
            res = fetch_robots(session, domain, args.timeout, args.www_fallback)
            fs = res["fetch_status"]

            if fs == "ok":
                robots_exists = True
                save_raw(raw_dir, domain, res["content"])
            elif fs == "not_found":
                robots_exists = False
            else:
                robots_exists = None

            if fs in ("ok", "not_found"):
                an = analyze(res["content"] or "", crawlers, domain)
                decisions = an["decisions"]
                explicit = an["explicit"]
                global_block = an["global_block_all"]
                cd = an["crawl_delay_wildcard"]
                smap = an["sitemap_count"]
            else:
                decisions = {c: "NA" for c in crawlers}
                explicit = {c: False for c in crawlers}
                global_block = ""
                cd = ""
                smap = ""

            row = {
                "no": site["no"],
                "layer": site["layer"],
                "name": site["name"],
                "domain": domain,
                "type": site["type"],
                "fetched_at": fetched_at,
                "fetch_status": fs,
                "http_status": res["http_status"],
                "final_url": res["final_url"],
                "robots_exists": "" if robots_exists is None else robots_exists,
                "robots_bytes": len(res["content"]) if res["content"] is not None else "",
                "global_block_all": global_block,
                "crawl_delay_wildcard": "" if cd is None else cd,
                "sitemap_count": smap,
                "elapsed_sec": res["elapsed_sec"],
                "error_message": res["error"],
            }
            for c in crawlers:
                row["block__" + c] = decisions[c]
                row["explicit__" + c] = "Y" if explicit[c] else ""
            ww.writerow(row)

            for c in crawlers:
                dec = decisions[c]
                blocked = "" if dec == "NA" else (1 if dec == "BLOCK" else 0)
                lw.writerow(
                    [
                        domain,
                        site["layer"],
                        site["name"],
                        site["type"],
                        c,
                        dec,
                        blocked,
                        1 if explicit[c] else 0,
                        "" if robots_exists is None else robots_exists,
                        fs,
                    ]
                )

            wf.flush()
            lf.flush()

            gb = "GLOBALBLOCK" if global_block is True else ""
            print(
                "[{}/{}] {:<28} {:<16} http={} {}".format(
                    i, n, domain, fs, res["http_status"], gb
                ).rstrip()
            )

            if i < n:
                time.sleep(random.uniform(args.delay_min, args.delay_max))

    summary = write_summary(wide_path, long_path, raw_dir, summary_path, stamp)
    print("\n" + summary)
    print("\n요약 파일:", summary_path)


if __name__ == "__main__":
    main()
