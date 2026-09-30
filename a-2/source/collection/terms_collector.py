import argparse
import csv
import os
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path

from html.parser import HTMLParser
from urllib.parse import urljoin, urlparse

try:
    import requests
except ImportError:
    sys.exit("[의존성 오류] requests 미설치 -> pip install requests")
try:
    import openpyxl
except ImportError:
    sys.exit("[의존성 오류] openpyxl 미설치 -> pip install openpyxl")


INPUTS = Path(__file__).resolve().parents[1] / "data" / "inputs"

BROWSER_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 (KHTML, like Gecko) "
    "Chrome/126.0.0.0 Safari/537.36"
)


ANCHOR_STRONG = [
    "이용약관",
    "서비스약관",
    "서비스 약관",
    "이용규정",
    "terms of service",
    "terms of use",
    "이용 약관",
]
ANCHOR_WEAK = ["약관", "terms", "법적고지", "법적 고지"]


HREF_HINTS = [
    "terms",
    "약관",
    "agreement",
    "stipulation",
    "policy/service",
    "rule",
    "clause",
    "provision",
    "member/terms",
    "policy/terms",
]


FALLBACK_PATHS = [
    "/terms",
    "/policy",
    "/agreement",
    "/terms-of-service",
    "/terms.html",
    "/policy/terms",
    "/agreement/terms",
    "/member/terms",
    "/help/terms",
]


KEYWORDS = [
    ("AI", re.compile(r"(?<![A-Za-z])AI(?![A-Za-z])")),
    ("인공지능", re.compile(r"인공지능")),
    ("기계학습", re.compile(r"기계학습")),
    ("머신러닝", re.compile(r"머신\s*러닝")),
    ("딥러닝", re.compile(r"딥\s*러닝")),
    ("학습데이터", re.compile(r"학습\s*데이터")),
    ("데이터마이닝", re.compile(r"데이터\s*마이닝")),
    ("텍스트마이닝", re.compile(r"텍스트\s*마이닝")),
    ("크롤링", re.compile(r"크롤링")),
    ("크롤러", re.compile(r"크롤러")),
    ("스크래핑", re.compile(r"스크래핑|스크레이핑")),
    ("무단수집", re.compile(r"무단\s*수집")),
    ("자동화된 수단", re.compile(r"자동화된\s*수단")),
    ("로봇", re.compile(r"로봇")),
    ("봇", re.compile(r"봇")),
]


TERMS_PAGE_SIGNAL = re.compile(r"약관|이용규정|terms", re.IGNORECASE)


class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.links = []
        self._footer_stack = []
        self._depth = 0
        self._footer_depths = []
        self._cur = None

    def handle_starttag(self, tag, attrs):
        ad = {k.lower(): (v or "") for k, v in attrs}

        is_footer = tag == "footer"
        blob = (ad.get("class", "") + " " + ad.get("id", "") + " " + ad.get("role", "")).lower()
        if "footer" in blob or ad.get("role", "").lower() == "contentinfo":
            is_footer = True
        self._depth += 1
        if is_footer:
            self._footer_depths.append(self._depth)
        if tag == "a":
            href = ad.get("href", "")
            self._cur = {"href": href, "texts": [], "footer": len(self._footer_depths) > 0}

    def handle_endtag(self, tag):
        if tag == "a" and self._cur is not None:
            text = " ".join(t.strip() for t in self._cur["texts"] if t.strip())
            self.links.append((self._cur["href"], text.strip(), self._cur["footer"]))
            self._cur = None

        if self._footer_depths and self._depth in self._footer_depths:
            self._footer_depths = [d for d in self._footer_depths if d != self._depth]
        self._depth = max(0, self._depth - 1)

    def handle_data(self, data):
        if self._cur is not None:
            self._cur["texts"].append(data)


class TextParser(HTMLParser):
    BLOCK = {
        "p",
        "br",
        "div",
        "li",
        "tr",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "section",
        "article",
        "header",
        "footer",
        "ul",
        "ol",
        "table",
        "td",
        "th",
        "dd",
        "dt",
        "blockquote",
        "pre",
    }
    SKIP = {"script", "style", "noscript", "head", "svg"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self._skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip_depth += 1
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_startendtag(self, tag, attrs):
        if tag == "br":
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip_depth > 0:
            self._skip_depth -= 1
        if tag in self.BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if self._skip_depth == 0:
            self.parts.append(data)

    def get_text(self):
        raw = "".join(self.parts)

        raw = re.sub(r"[ \t 　]+", " ", raw)
        raw = re.sub(r"\n[ \t]+", "\n", raw)
        raw = re.sub(r"[ \t]+\n", "\n", raw)
        raw = re.sub(r"\n{2,}", "\n", raw)
        return raw.strip()


def bare_www_candidates(domain):

    hosts = [domain]
    if domain.startswith("www."):
        return hosts
    labels = domain.split(".")
    kr_second = {"co", "or", "go", "ac", "re", "ne", "kr", "pe"}
    is_bare = (len(labels) == 2) or (len(labels) == 3 and labels[1] in kr_second)
    if is_bare:
        hosts.append("www." + domain)
    return hosts


def decode_response(resp):

    ctype = resp.headers.get("Content-Type", "").lower()
    enc = None
    if "charset=" in ctype:
        enc = resp.encoding
    if not enc:
        enc = resp.apparent_encoding or "utf-8"
    try:
        return resp.content.decode(enc, errors="replace")
    except (LookupError, TypeError):
        return resp.content.decode("utf-8", errors="replace")


def fetch(session, url, timeout):

    r = {"http_status": None, "final_url": None, "text": None, "ok": False, "error": ""}
    try:
        resp = session.get(url, timeout=timeout, allow_redirects=True)
        r["http_status"] = resp.status_code
        r["final_url"] = resp.url
        ctype = resp.headers.get("Content-Type", "").lower()
        if resp.status_code == 200 and ("html" in ctype or "text" in ctype or not ctype):
            r["text"] = decode_response(resp)
            r["ok"] = True
        elif resp.status_code == 200:
            r["error"] = "non-text content-type: " + ctype
        else:
            r["error"] = "HTTP {}".format(resp.status_code)
    except requests.exceptions.Timeout as e:
        r["error"] = "timeout: " + str(e)[:120]
    except requests.exceptions.SSLError as e:
        r["error"] = "ssl: " + str(e)[:120]
    except requests.exceptions.ConnectionError as e:
        r["error"] = "conn: " + str(e)[:120]
    except requests.exceptions.RequestException as e:
        r["error"] = "req: " + str(e)[:120]
    except Exception as e:
        r["error"] = "other: " + str(e)[:120]
    return r


def get_homepage(session, domain, timeout):

    last = None
    for host in bare_www_candidates(domain):
        res = fetch(session, "https://{}/".format(host), timeout)
        last = res
        if res["ok"] or res["http_status"] is not None:
            return res
    return last


def score_link(href, text, in_footer):

    if not href:
        return -1
    low_href = href.lower()
    if low_href.startswith(("javascript:", "mailto:", "tel:", "#")):
        return -1
    t = (text or "").lower().strip()
    score = 0

    strong = any(a in t for a in ANCHOR_STRONG)
    weak = any(a in t for a in ANCHOR_WEAK)
    if strong:
        score += 100
    elif weak:
        score += 60

    if any(h in low_href for h in HREF_HINTS):
        score += 25

    if in_footer:
        score += 20

    if not strong and not weak:
        if "privacy" in low_href or "개인정보" in t:
            return -1

        if score <= 25:
            score -= 40
    return score


def find_terms_link(html, base_url):

    try:
        p = LinkParser()
        p.feed(html)
    except Exception:
        return None
    best = None
    best_score = 0
    for href, text, in_footer in p.links:
        s = score_link(href, text, in_footer)
        if s > best_score:
            best_score = s
            best = urljoin(base_url, href)

    if best_score >= 60:
        return best
    return None


def split_sentences(text):

    tmp = re.sub(r"([.!?。])(\s)", r"\1\n", text)
    sents = [s.strip() for s in re.split(r"[\r\n]+", tmp) if s.strip()]
    return sents


def find_keywords_in(s):

    found = []
    for name, rx in KEYWORDS:
        if rx.search(s):
            found.append(name)
    return found


def extract_contexts(text, max_chars=20000):

    sents = split_sentences(text)
    match_idx = {}
    for i, s in enumerate(sents):
        kws = find_keywords_in(s)
        if kws:
            match_idx[i] = kws

    if not match_idx:
        return [], []

    keep = set()
    for i in match_idx:
        for j in (i - 1, i, i + 1):
            if 0 <= j < len(sents):
                keep.add(j)

    blocks = []
    for idx in sorted(keep):
        if blocks and idx == blocks[-1][-1] + 1:
            blocks[-1].append(idx)
        else:
            blocks.append([idx])

    out_blocks = []
    total = 0
    for blk in blocks:
        chunk = " ".join(sents[k] for k in blk).strip()
        if not chunk:
            continue
        if total + len(chunk) > max_chars:
            out_blocks.append("[...이하 생략: 추출 길이 상한 도달...]")
            break
        out_blocks.append(chunk)
        total += len(chunk)

    found_union = []
    seen = set()
    for i in sorted(match_idx):
        for k in match_idx[i]:
            if k not in seen:
                seen.add(k)
                found_union.append(k)
    order = {name: n for n, (name, _) in enumerate(KEYWORDS)}
    found_union.sort(key=lambda k: order.get(k, 999))
    return out_blocks, found_union


def process_domain(session, domain, timeout, raw_dir):

    now = datetime.now().isoformat(timespec="seconds")
    result = {
        "domain": domain,
        "terms_url": "",
        "status": "접속불가",
        "keywords": "",
        "sentences": "",
        "raw_len": 0,
        "collected_at": now,
        "note": "",
    }

    home = get_homepage(session, domain, timeout)
    if home is None or not home["ok"] or not home["text"]:
        result["status"] = "접속불가"
        result["note"] = (home or {}).get("error", "no homepage")
        return result

    base_url = home["final_url"] or ("https://%s/" % domain)

    terms_url = find_terms_link(home["text"], base_url)
    terms_res = None
    if terms_url:
        terms_res = fetch(session, terms_url, timeout)
        if not terms_res["ok"] or not terms_res["text"]:
            terms_res = None

    if terms_res is None or not terms_res.get("ok"):
        root = "{u.scheme}://{u.netloc}".format(u=urlparse(base_url))
        for path in FALLBACK_PATHS:
            cand = root + path
            cres = fetch(session, cand, timeout)
            if cres["ok"] and cres["text"] and TERMS_PAGE_SIGNAL.search(cres["text"][:6000]):
                terms_url = cres["final_url"] or cand
                terms_res = cres
                result["note"] = "관용경로:" + path
                break
            time.sleep(0.3)

    if terms_res is None or not terms_res.get("ok") or not terms_res.get("text"):
        result["status"] = "링크못찾음"
        return result

    parser = TextParser()
    try:
        parser.feed(terms_res["text"])
        page_text = parser.get_text()
    except Exception as e:
        page_text = ""
        result["note"] = (result["note"] + " | textparse:" + str(e)[:60]).strip(" |")

    result["terms_url"] = terms_url
    result["status"] = "성공"
    result["raw_len"] = len(page_text)

    if page_text:
        safe = re.sub(r"[^0-9A-Za-z._-]", "_", domain)
        with open(os.path.join(raw_dir, safe + ".txt"), "w", encoding="utf-8", newline="") as f:
            f.write(
                "# domain: %s\n# terms_url: %s\n# collected_at: %s\n\n" % (domain, terms_url, now)
            )
            f.write(page_text)

    blocks, found = extract_contexts(page_text)
    if found:
        result["keywords"] = ", ".join(found)
        result["sentences"] = "\n\n---\n\n".join(blocks)
    else:
        result["keywords"] = "키워드없음"
        result["sentences"] = "키워드없음"
    return result


def load_sites(xlsx_path):

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["서브샘플_181"]
    sites = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        no = row[0]
        if not isinstance(no, (int, float)):
            continue
        domain = row[3]
        if not domain or not str(domain).strip():
            continue
        sites.append(
            {
                "no": int(no),
                "layer": str(row[1]).strip() if row[1] else "",
                "name": str(row[2]).strip() if row[2] else "",
                "domain": str(domain).strip().lower(),
                "type": str(row[4]).strip() if row[4] else "",
            }
        )
    wb.close()
    return sites


def write_xlsx_copy(src_xlsx, out_xlsx, results_by_domain):

    wb = openpyxl.load_workbook(src_xlsx)
    ws = wb["서브샘플_181"]

    hmap = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v is not None:
            hmap[str(v).strip()] = c
    col_url = hmap.get("약관URL", 12)
    col_stat = hmap.get("탐색성공여부", 13)
    col_kw = hmap.get("발견키워드", 14)
    col_sent = hmap.get("추출문장", 15)

    filled = 0
    for r in range(2, ws.max_row + 1):
        dom = ws.cell(r, 4).value
        if not dom:
            continue
        key = str(dom).strip().lower()
        res = results_by_domain.get(key)
        if not res:
            continue
        ws.cell(r, col_url).value = res["terms_url"]
        ws.cell(r, col_stat).value = res["status"]
        ws.cell(r, col_kw).value = res["keywords"]
        ws.cell(r, col_sent).value = res["sentences"]
        filled += 1
    wb.save(out_xlsx)
    return filled


def write_summary(results, sites, summary_path, stamp):
    n = len(results)
    from collections import Counter

    status_ct = Counter(r["status"] for r in results)
    kw_found = sum(
        1 for r in results if r["status"] == "성공" and r["keywords"] not in ("", "키워드없음")
    )
    success = status_ct.get("성공", 0)

    kw_domain_ct = Counter()
    for r in results:
        if r["keywords"] and r["keywords"] != "키워드없음":
            for k in [x.strip() for x in r["keywords"].split(",")]:
                if k:
                    kw_domain_ct[k] += 1

    layer_of = {s["domain"]: s["layer"] for s in sites}
    layer_total = Counter()
    layer_success = Counter()
    layer_kw = Counter()
    for r in results:
        L = layer_of.get(r["domain"], "(미지정)")
        layer_total[L] += 1
        if r["status"] == "성공":
            layer_success[L] += 1
        if r["keywords"] and r["keywords"] != "키워드없음":
            layer_kw[L] += 1

    lines = []
    lines.append("=== RQ3 이용약관 자동수집 요약 (%s) ===" % stamp)
    lines.append("대상 도메인: %d" % n)
    lines.append("")
    lines.append("탐색성공여부 분포:")
    for k in ("성공", "링크못찾음", "접속불가"):
        v = status_ct.get(k, 0)
        lines.append("    %-10s %4d  (%.1f%%)" % (k, v, 100.0 * v / n if n else 0))
    for k in status_ct:
        if k not in ("성공", "링크못찾음", "접속불가"):
            lines.append("    %-10s %4d" % (k, status_ct[k]))
    lines.append("")
    lines.append(
        "키워드 발견율(성공 도메인 기준): %d/%d = %.1f%%"
        % (kw_found, success, 100.0 * kw_found / success if success else 0)
    )
    lines.append(
        "키워드 발견율(전체 도메인 기준): %d/%d = %.1f%%"
        % (kw_found, n, 100.0 * kw_found / n if n else 0)
    )
    lines.append("")
    lines.append("키워드별 발견 도메인 수(다중 카운트):")
    if kw_domain_ct:
        order = {name: i for i, (name, _) in enumerate(KEYWORDS)}
        for k in sorted(kw_domain_ct, key=lambda x: order.get(x, 999)):
            lines.append("    %-14s %3d" % (k, kw_domain_ct[k]))
    else:
        lines.append("    (없음)")
    lines.append("")
    lines.append("층별 요약 (건수 / 탐색성공 / 키워드발견):")
    for L in sorted(layer_total):
        lines.append(
            "    %-14s 건수 %3d  탐색성공 %3d (%.0f%%)  키워드발견 %3d"
            % (
                L,
                layer_total[L],
                layer_success[L],
                100.0 * layer_success[L] / layer_total[L] if layer_total[L] else 0,
                layer_kw[L],
            )
        )
    summary = "\n".join(lines)
    with open(summary_path, "w", encoding="utf-8", newline="") as f:
        f.write(summary + "\n")
    return summary


CSV_COLS = [
    "도메인",
    "약관URL",
    "탐색성공여부",
    "발견키워드",
    "추출문장",
    "원문길이(chars)",
    "수집일시",
]


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    ap = argparse.ArgumentParser(description="이용약관 자동 탐색·추출 수집기(RQ3)")
    ap.add_argument("--xlsx", default=str(INPUTS / "terms_sample.xlsx"))
    ap.add_argument("--outdir", required=True, help="새 실행의 출력 폴더(기존 폴더 사용 불가)")
    ap.add_argument("--limit", type=int, default=None, help="앞에서 N개만")
    ap.add_argument("--start", type=int, default=0, help="건너뛸 개수(오프셋)")
    ap.add_argument("--timeout", type=float, default=10.0)
    ap.add_argument("--delay-min", type=float, default=1.0)
    ap.add_argument("--delay-max", type=float, default=2.0)
    ap.add_argument("--write-xlsx", action="store_true", help="수집 후 L~O열 채운 엑셀 사본 저장")
    ap.add_argument(
        "--xlsx-out", default=None, help="새 엑셀 사본 경로(기본: 출력 폴더의 terms_collected.xlsx)"
    )
    args = ap.parse_args()
    if args.timeout <= 0 or args.delay_min < 0 or args.delay_max < args.delay_min:
        ap.error("timeout > 0 및 0 <= delay-min <= delay-max 조건이 필요합니다.")
    if args.start < 0 or (args.limit is not None and args.limit < 0):
        ap.error("start, limit는 음수일 수 없습니다.")
    if Path(args.outdir).exists():
        ap.error("출력 경로가 이미 존재합니다. 새로운 실행 폴더를 지정하세요.")
    if args.write_xlsx:
        args.xlsx_out = args.xlsx_out or str(Path(args.outdir) / "terms_collected.xlsx")
        if Path(args.xlsx_out).exists():
            ap.error("엑셀 출력 경로가 이미 존재합니다. 새 사본 경로를 지정하세요.")
        if Path(args.xlsx_out).resolve() == Path(args.xlsx).resolve():
            ap.error("엑셀 출력 경로는 입력 파일과 달라야 합니다.")

    if not os.path.exists(args.xlsx):
        sys.exit("[오류] 엑셀 없음: %s" % args.xlsx)

    sites = load_sites(args.xlsx)
    sel = sites[args.start :]
    if args.limit is not None:
        sel = sel[: args.limit]

    os.makedirs(args.outdir, exist_ok=False)
    raw_dir = os.path.join(args.outdir, "terms_raw")
    os.makedirs(raw_dir, exist_ok=True)
    csv_path = os.path.join(args.outdir, "terms_extraction_results.csv")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    summary_path = os.path.join(args.outdir, "run_summary.txt")

    print(
        "입력 %d개 도메인 중 이번 실행 대상 %d개 -> outdir='%s'"
        % (len(sites), len(sel), args.outdir)
    )

    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": BROWSER_UA,
            "Accept": ("text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"),
            "Accept-Language": "ko,en;q=0.8",
        }
    )

    results = []
    n = len(sel)
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as cf:
        w = csv.DictWriter(cf, fieldnames=CSV_COLS)
        w.writeheader()
        for i, site in enumerate(sel, 1):
            domain = site["domain"]
            try:
                res = process_domain(session, domain, args.timeout, raw_dir)
            except Exception as e:
                res = {
                    "domain": domain,
                    "terms_url": "",
                    "status": "접속불가",
                    "keywords": "",
                    "sentences": "",
                    "raw_len": 0,
                    "collected_at": datetime.now().isoformat(timespec="seconds"),
                    "note": "예외:" + str(e)[:100],
                }
            results.append(res)

            w.writerow(
                {
                    "도메인": res["domain"],
                    "약관URL": res["terms_url"],
                    "탐색성공여부": res["status"],
                    "발견키워드": res["keywords"],
                    "추출문장": res["sentences"],
                    "원문길이(chars)": res["raw_len"],
                    "수집일시": res["collected_at"],
                }
            )
            cf.flush()

            kwshow = res["keywords"] if res["keywords"] else "-"
            print("[%d/%d] %-28s %-8s kw=%s" % (i, n, domain, res["status"], kwshow[:40]))

            if i < n:
                time.sleep(random.uniform(args.delay_min, args.delay_max))

    summary = write_summary(results, sites, summary_path, stamp)
    print("\n" + summary)
    print("\nCSV:", csv_path)
    print("요약:", summary_path)

    if args.write_xlsx:
        results_by_domain = {r["domain"]: r for r in results}
        out_xlsx = args.xlsx_out
        Path(out_xlsx).parent.mkdir(parents=True, exist_ok=True)
        filled = write_xlsx_copy(args.xlsx, out_xlsx, results_by_domain)
        print("엑셀 사본 저장: %s  (L~O 채운 행 %d)" % (out_xlsx, filled))


if __name__ == "__main__":
    main()
